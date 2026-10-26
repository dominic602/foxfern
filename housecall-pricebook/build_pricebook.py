"""Build the Fox & Fern Housecall Pro labor-rate price book (CSV) and rate calculator (XLSX).

Edit the rates below, run `python3 build_pricebook.py`, and re-import the CSV.

  Labor cost/hr = sub crew rate + PM loaded hourly x PM hours per field hour
  GP %          = (sell rate - labor cost) / sell rate, must be >= MIN_GP
Products/materials are uploaded separately in Housecall Pro.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

OUT = Path(__file__).parent

PM_WEEKLY_PAY = 1000.00   # PM gross pay per week
PM_BURDEN = 0.25          # payroll taxes, workers comp, etc.
PM_WEEKLY_HOURS = 40
MIN_GP = 0.33
NET30_PCT = 0.05
RES_MAT_MARKUP = 1.5
COM_MAT_MARKUP = 2.0

PM_HOURLY = PM_WEEKLY_PAY * (1 + PM_BURDEN) / PM_WEEKLY_HOURS

# (category, name, sub cost/hr, PM hrs per field hr, sell rate/hr, description)
LINES = [
    ("Residential Remodeling", "Residential Interior Remodeling Labor", 60, 0.25, 115,
     "Interior remodeling labor per hour: demo, framing, drywall, trim, cabinets, tile, flooring. "
     "Includes project management. Qty = labor hours (1 day = 8 hrs)."),
    ("Residential Remodeling", "Residential Exterior Remodeling Labor", 63, 0.25, 120,
     "Exterior remodeling labor per hour: decks, fencing, siding, trim/fascia, doors and windows. "
     "Includes project management. Qty = labor hours (1 day = 8 hrs)."),
    ("Residential Remodeling", "Residential Handyman Labor", 55, 0.10, 115,
     "Handyman labor per hour, 2-hour minimum per visit. Qty = labor hours."),
    ("Residential Remodeling", "Residential Concrete & Hardscape Labor", 70, 0.25, 130,
     "Concrete, pavers, retaining walls and grading labor per hour. "
     "Includes project management. Qty = labor hours (1 day = 8 hrs)."),
    ("Residential Remodeling", "Residential Demolition & Haul-Away Labor", 48, 0.25, 95,
     "Demolition, site protection, debris loading and haul-away labor per hour. "
     "Dump/dumpster fees billed separately. Qty = labor hours."),
    ("Commercial Remodeling", "Commercial Interior Remodeling Labor", 65, 0.30, 150,
     "Commercial interior remodeling labor per hour, Montgomery County and Philadelphia. "
     "Includes project management, COI and site coordination. Qty = labor hours."),
    ("Commercial Remodeling", "Commercial Exterior Remodeling Labor", 68, 0.30, 155,
     "Commercial exterior remodeling labor per hour, Montgomery County and Philadelphia. "
     "Includes project management, COI and site coordination. Qty = labor hours."),
    ("Commercial Remodeling", "Commercial Handyman / Facility Maintenance Labor", 60, 0.30, 140,
     "Commercial handyman and facility maintenance labor per hour, 2-hour minimum. Qty = labor hours."),
    ("Commercial Remodeling", "Commercial Concrete & Hardscape Labor", 75, 0.30, 165,
     "Commercial concrete, sidewalk, pad and hardscape labor per hour. "
     "Includes project management. Qty = labor hours."),
]
NET30 = ("Payment Terms", "Net 30 Payment Terms Adjustment (per $100)", NET30_PCT * 100,
         "Net 30 invoice terms add 5% to the contract. Qty = invoice subtotal / 100 "
         "(e.g. $24,000 = qty 240).")

# Housecall Pro price book (services) columns. If your exported file's header row
# differs, change only these labels -- the data order stays the same.
HEADER = ["Category", "Subcategory", "Name", "Description", "Price", "Cost", "Unit of measure"]
CATEGORY = "Construction & Remodeling"


def labor_cost(sub, pm_ratio):
    return sub + PM_HOURLY * pm_ratio


def write_csv():
    path = OUT / "foxfern_housecall_pricebook.csv"
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for cat, name, sub, pm, rate, desc in LINES:
            w.writerow([CATEGORY, cat, name, desc, f"{rate:.2f}", f"{labor_cost(sub, pm):.2f}", "Hour"])
        cat, name, price, desc = NET30
        w.writerow([CATEGORY, cat, name, desc, f"{price:.2f}", "0.00", "Each"])
    return path


def write_xlsx():
    wb = Workbook()
    inp = PatternFill("solid", fgColor="FFF2CC")
    head_fill = PatternFill("solid", fgColor="2F4F3A")
    head_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "Labor Rates"
    ws["A1"] = "Fox & Fern - Labor Rates (yellow cells are editable)"
    ws["A1"].font = Font(bold=True, size=14)
    settings = [
        ("PM weekly pay", PM_WEEKLY_PAY, "$#,##0.00", "PM_Pay"),
        ("PM payroll burden %", PM_BURDEN, "0%", "PM_Burden"),
        ("PM hours per week", PM_WEEKLY_HOURS, "0", "PM_Hours"),
        ("Minimum gross profit %", MIN_GP, "0%", "Min_GP"),
        ("Residential material markup (x)", RES_MAT_MARKUP, "0.00", "RES_Mat"),
        ("Commercial material markup (x)", COM_MAT_MARKUP, "0.00", "COM_Mat"),
        ("Net 30 adjustment %", NET30_PCT, "0%", "Net30"),
    ]
    for i, (label, val, fmt, nm) in enumerate(settings, start=3):
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, val)
        c.fill, c.number_format = inp, fmt
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"'Labor Rates'!$B${i}")
    ws.cell(10, 1, "PM loaded cost / hr")
    ws.cell(10, 2, "=PM_Pay*(1+PM_Burden)/PM_Hours").number_format = "$#,##0.00"
    wb.defined_names["PM_Hourly"] = DefinedName("PM_Hourly", attr_text="'Labor Rates'!$B$10")

    cols = ["Labor line", "Sub cost / hr", "PM hrs per field hr", "Sell rate / hr",
            "Loaded cost / hr", "GP %", "Min rate for GP", "Day rate (8 hrs)", "Check"]
    top = 12
    for j, h in enumerate(cols, start=1):
        c = ws.cell(top, j, h)
        c.font, c.fill = head_font, head_fill
    for i, (cat, name, sub, pm, rate, _) in enumerate(LINES, start=top + 1):
        ws.cell(i, 1, name)
        for j, (v, fmt) in enumerate(((sub, "$#,##0.00"), (pm, "0.00"), (rate, "$#,##0.00")), start=2):
            c = ws.cell(i, j, v)
            c.fill, c.number_format = inp, fmt
        ws.cell(i, 5, f"=B{i}+PM_Hourly*C{i}").number_format = "$#,##0.00"
        ws.cell(i, 6, f"=(D{i}-E{i})/D{i}").number_format = "0.0%"
        ws.cell(i, 7, f"=E{i}/(1-Min_GP)").number_format = "$#,##0.00"
        ws.cell(i, 8, f"=D{i}*8").number_format = "$#,##0"
        ws.cell(i, 9, f'=IF(F{i}<Min_GP,"RAISE RATE","OK")')
    for col, w in zip("ABCDEFGHI", (46, 13, 12, 13, 14, 9, 14, 14, 12)):
        ws.column_dimensions[col].width = w

    # Job estimator: hours per labor line + material cost
    je = wb.create_sheet("Job Estimator")
    je["A1"] = "Job Estimator - enter labor hours and material cost"
    je["A1"].font = Font(bold=True, size=14)
    for j, h in enumerate(["Labor line", "Hours", "Price", "Cost"], start=1):
        c = je.cell(3, j, h)
        c.font, c.fill = head_font, head_fill
    first = 4
    for k, (_, name, *_rest) in enumerate(LINES):
        i = first + k
        r = top + 1 + k
        je.cell(i, 1, name)
        c = je.cell(i, 2, 0)
        c.fill = inp
        je.cell(i, 3, f"=B{i}*'Labor Rates'!D{r}").number_format = "$#,##0"
        je.cell(i, 4, f"=B{i}*'Labor Rates'!E{r}").number_format = "$#,##0"
    last = first + len(LINES) - 1
    r = last + 2
    rows = [
        ("Commercial job? (Y/N)", "N", None, None),
        ("Material cost ($, your cost)", 0, '=B{m}*IF(B{c}="Y",COM_Mat,RES_Mat)', "=B{m}"),
        ("Net 30? (Y/N)", "N", None, None),
        ("Subtotal", None, f"=SUM(C{first}:C{last})+C{{m}}", f"=SUM(D{first}:D{last})+D{{m}}"),
        ("Net 30 adjustment", None, '=IF(B{n}="Y",C{s}*Net30,0)', None),
        ("TOTAL PRICE", None, "=C{s}+C{a}", "=D{s}"),
        ("Gross profit %", None, "=IF(C{t}=0,0,(C{t}-D{t})/C{t})", None),
        ("Material cost % of price (target <= 20%)", None, "=IF(C{t}=0,0,B{m}/C{t})", None),
    ]
    idx = dict(c=r, m=r + 1, n=r + 2, s=r + 3, a=r + 4, t=r + 5)
    for k, (label, val, pf, cf) in enumerate(rows):
        i = r + k
        je.cell(i, 1, label)
        if val is not None:
            c = je.cell(i, 2, val)
            c.fill = inp
            if isinstance(val, (int, float)):
                c.number_format = "$#,##0.00"
        if pf:
            c = je.cell(i, 3, pf.format(**idx))
            c.number_format = "0.0%" if "%" in label else "$#,##0"
        if cf:
            je.cell(i, 4, cf.format(**idx)).number_format = "$#,##0"
    je.cell(r + 5, 1).font = Font(bold=True)
    je.cell(r + 5, 3).font = Font(bold=True)
    je.column_dimensions["A"].width = 46
    for col in "BCD":
        je.column_dimensions[col].width = 13

    path = OUT / "foxfern_pricing_calculator.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"PM loaded cost/hr: ${PM_HOURLY:.2f}")
    low = []
    for _, name, sub, pm, rate, _ in LINES:
        cost = labor_cost(sub, pm)
        gp = (rate - cost) / rate
        print(f"  {name:50s} ${rate:7.2f}/hr  cost ${cost:6.2f}  GP {gp:.1%}")
        if gp < MIN_GP:
            low.append(name)
    print(f"-> {write_csv().name}, {write_xlsx().name}")
    if low:
        raise SystemExit("Below minimum GP: " + ", ".join(low))
